import zipfile, io
from pathlib import Path
from docx import Document
from PIL import Image
import pytesseract

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def ocr_image(img):
    img = img.convert("L")  # grayscale
    img = img.resize((img.width * 2, img.height * 2))  # 2x bada
    return pytesseract.image_to_string(img, config="--psm 6")


def load_image(path):
    return ocr_image(Image.open(path))


def docx_embedded_images_text(path):
    """OCR any pictures stored inside the docx (scanned pages)."""
    out = []
    with zipfile.ZipFile(path) as z:
        for name in sorted(z.namelist()):
            if name.startswith("word/media/") and name.lower().endswith(
                (".png", ".jpg", ".jpeg")
            ):
                out.append(ocr_image(Image.open(io.BytesIO(z.read(name)))))
    return "\n".join(out)


def load_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(parts)

    img_text = docx_embedded_images_text(path)
    if img_text.strip():
        text += (
            "\n\n[OCR OF EMBEDDED PAGE IMAGES - may be cleaner than the text above]\n"
            + img_text
        )
    return text


def load_text(path):
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return load_docx(path)
    if ext in (".png", ".jpg", ".jpeg"):
        return load_image(path)
    raise ValueError(f"Unsupported file type: {path}")
